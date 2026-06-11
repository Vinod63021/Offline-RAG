import os
import json
import sqlite3
import pandas as pd
import easyocr

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pdf2image import convert_from_path

from langchain.schema import Document
from langchain_community.document_loaders import PyPDFLoader


# ==========================================
# OCR INITIALIZATION
# ==========================================

print("Initializing OCR...")

ocr_reader = easyocr.Reader(
    ['en'],
    gpu=False
)


# ==========================================
# SCANNED PDF OCR
# ==========================================

def load_scanned_pdf(path):

    pages = convert_from_path(path)

    full_text = ""

    for i, page in enumerate(pages):

        temp_path = f"temp_page_{i}.png"

        page.save(
            temp_path,
            "PNG"
        )

        text = "\n".join(
            ocr_reader.readtext(
                temp_path,
                detail=0
            )
        )

        full_text += text + "\n"

        if os.path.exists(temp_path):
            os.remove(temp_path)

    return [
        Document(
            page_content=full_text,
            metadata={
                "source": path,
                "type": "scanned_pdf"
            }
        )
    ]


# ==========================================
# PDF
# ==========================================

def load_pdf(path):

    try:

        loader = PyPDFLoader(path)

        docs = loader.load()

        extracted_text = ""

        for doc in docs:

            extracted_text += (
                doc.page_content.strip()
            )

        if len(extracted_text) > 100:

            for doc in docs:

                doc.metadata["type"] = "pdf"

            return docs

        print(
            f"Scanned PDF detected: {os.path.basename(path)}"
        )

        return load_scanned_pdf(path)

    except Exception:

        print(
            f"OCR fallback for: {os.path.basename(path)}"
        )

        return load_scanned_pdf(path)


# ==========================================
# DOCX
# ==========================================

def load_docx(path):

    doc = DocxDocument(path)

    text = "\n".join(
        para.text
        for para in doc.paragraphs
    )

    return [
        Document(
            page_content=text,
            metadata={
                "source": path,
                "type": "docx"
            }
        )
    ]


# ==========================================
# EXCEL
# ==========================================

def load_excel(path):

    df = pd.read_excel(path)

    text = df.to_string(index=False)

    return [
        Document(
            page_content=text,
            metadata={
                "source": path,
                "type": "xlsx"
            }
        )
    ]


# ==========================================
# CSV
# ==========================================

def load_csv(path):

    df = pd.read_csv(path)

    text = df.to_string(index=False)

    return [
        Document(
            page_content=text,
            metadata={
                "source": path,
                "type": "csv"
            }
        )
    ]


# ==========================================
# TXT
# ==========================================

def load_txt(path):

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as f:

        text = f.read()

    return [
        Document(
            page_content=text,
            metadata={
                "source": path,
                "type": "txt"
            }
        )
    ]


# ==========================================
# JSON
# ==========================================

def load_json(path):

    with open(
        path,
        "r",
        encoding="utf-8"
    ) as f:

        data = json.load(f)

    text = json.dumps(
        data,
        indent=2,
        ensure_ascii=False
    )

    return [
        Document(
            page_content=text,
            metadata={
                "source": path,
                "type": "json"
            }
        )
    ]


# ==========================================
# IMAGE OCR
# ==========================================

def load_image(path):

    text = "\n".join(
        ocr_reader.readtext(
            path,
            detail=0
        )
    )

    return [
        Document(
            page_content=text,
            metadata={
                "source": path,
                "type": "image"
            }
        )
    ]


# ==========================================
# SQLITE DATABASE
# ==========================================

def load_sqlite(path):

    documents = []

    conn = sqlite3.connect(path)

    cursor = conn.cursor()

    tables = cursor.execute(
        """
        SELECT name
        FROM sqlite_master
        WHERE type='table';
        """
    ).fetchall()

    for table in tables:

        table_name = table[0]

        rows = cursor.execute(
            f"SELECT * FROM {table_name}"
        ).fetchall()

        text = f"Table: {table_name}\n\n"

        for row in rows:

            text += str(row) + "\n"

        documents.append(
            Document(
                page_content=text,
                metadata={
                    "source": path,
                    "table": table_name,
                    "type": "sqlite"
                }
            )
        )

    conn.close()

    return documents


# ==========================================
# HTML FILE
# ==========================================

def load_html(path):

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as f:

        html = f.read()

    soup = BeautifulSoup(
        html,
        "html.parser"
    )

    text = soup.get_text(
        separator="\n",
        strip=True
    )

    return [
        Document(
            page_content=text,
            metadata={
                "source": path,
                "type": "html"
            }
        )
    ]


# ==========================================
# MAIN LOADER
# ==========================================

def load_all_documents(data_folder):

    documents = []

    supported_extensions = {

        ".pdf": load_pdf,

        ".docx": load_docx,

        ".xlsx": load_excel,

        ".csv": load_csv,

        ".txt": load_txt,

        ".json": load_json,

        ".jpg": load_image,
        ".jpeg": load_image,
        ".png": load_image,
        ".bmp": load_image,
        ".webp": load_image,

        ".db": load_sqlite,
        ".sqlite": load_sqlite,

        ".html": load_html,
        ".htm": load_html
    }

    for root, dirs, files in os.walk(data_folder):

        for file in files:

            file_path = os.path.join(
                root,
                file
            )

            ext = os.path.splitext(
                file
            )[1].lower()

            if ext in supported_extensions:

                try:

                    docs = supported_extensions[
                        ext
                    ](file_path)

                    documents.extend(docs)

                    print(
                        f"Loaded: {file}"
                    )

                except Exception as e:

                    print(
                        f"Failed: {file}"
                    )

                    print(e)

    return documents